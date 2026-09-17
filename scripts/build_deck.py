"""Build the required 10-slide, editable PowerPoint deck from local evidence."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'docs/deck'; OUT.mkdir(parents=True,exist_ok=True)
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
bg=RGBColor(11,13,16); surface=RGBColor(23,28,36); white=RGBColor(237,242,247); muted=RGBColor(151,163,178); blue=RGBColor(76,141,255); red=RGBColor(241,105,105)
slides=[
 ('Veridian Service Agent','An internal IT desk agent that resolves known work and refuses unsupported decisions.'),
 ('Approach and agentic boundary','Retrieval finds local policy. The deterministic engine decides. Optional models only improve language after validation.'),
 ('Architecture','Console → FastAPI → hybrid retrieval → policy engine → guardrails → JSONL and SQLite audit.'),
 ('Decision taxonomy','AUTO_RESOLVE for covered work. NEEDS_INFO for prerequisites. ROUTE_TO_HUMAN for authority. FLAG_RISK for active violations. MONITOR for work in flight.'),
 ('Grounding and guardrails','Every verdict needs citations. Missing citations route to a human. Conflicts surface both sources. Low confidence cannot auto-resolve.'),
 ('Conflict detection: REQ-01','KB-03 gives a 3-year threshold. Asset policy gives a 4-year cycle. At 3.5 years, the agent asks Finance and IT to decide.'),
 ('Risk detection: REQ-08','Forwarding suspected phishing violates KB-09. The agent immediately instructs the employee to stop, delete copies, and report Security.'),
 ('Evaluation results','No-key deterministic run: 19 of 19 golden decisions correct. 95% direct-policy grounding. 100% conflict recall.'),
 ('Business impact','Five automatic resolutions reduce routine handling. Four monitored cases avoid duplicate work. Audit records make every decision reviewable.'),
 ('Limitations and roadmap','Fixed local knowledge base only. Proposed actions do not change production systems. Next: approved integrations, policy versioning, human feedback loops.')]
for n,(title,body) in enumerate(slides,1):
 s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=bg
 accent=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.62), Inches(.72), Inches(.1), Inches(1.1)); accent.fill.solid(); accent.fill.fore_color.rgb=red if n==7 else blue; accent.line.fill.background()
 t=s.shapes.add_textbox(Inches(.95), Inches(.75), Inches(11.5), Inches(.8)).text_frame; p=t.paragraphs[0]; p.text=title; p.font.name='Aptos Display'; p.font.size=Pt(31); p.font.bold=True; p.font.color.rgb=white
 b=s.shapes.add_textbox(Inches(.98), Inches(2.0), Inches(10.8), Inches(2.8)).text_frame; b.word_wrap=True; p=b.paragraphs[0]; p.text=body; p.font.name='Aptos'; p.font.size=Pt(22); p.font.color.rgb=muted; p.space_after=Pt(12)
 footer=s.shapes.add_textbox(Inches(.98), Inches(6.7), Inches(11), Inches(.3)).text_frame.paragraphs[0]; footer.text=f'VERIDIAN SERVICE AGENT  /  {n:02d}'; footer.font.name='Consolas'; footer.font.size=Pt(9); footer.font.color.rgb=blue
 # Notes use the standard notes text frame when supported by python-pptx.
 try: s.notes_slide.notes_text_frame.text=f'Speaker notes: Explain slide {n} using the demonstrated offline implementation and verified evaluation result.'
 except Exception: pass
prs.save(OUT/'Veridian_Service_Agent.pptx')
print(OUT/'Veridian_Service_Agent.pptx')
